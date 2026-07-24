"""교안 네이티브 렌더러 (v7) 테스트."""
from __future__ import annotations

import json

import pytest
from pptx import Presentation

from pptm.design.gyoan import render_deck, SLIDE_W, SLIDE_H


@pytest.fixture()
def project(tmp_path):
    pdir = tmp_path / "pages" / "p01"
    pdir.mkdir(parents=True)
    spec = {
        "page": 1, "variant": "paper",
        "items": [
            {"type": "label", "text": "Ⅰ. 추진 목적", "x": 0.05, "y": 0.06},
            {"type": "headline", "lines": ["한 줄 헤드라인입니다."], "x": 0.05, "y": 0.11, "w": 0.9},
            {"type": "card", "x": 0.05, "y": 0.3, "w": 0.9, "h": 0.12},
            {"type": "text", "text": "본문입니다", "x": 0.08, "y": 0.33, "w": 0.5, "style": "body"},
            {"type": "pill", "text": "150분", "x": 0.3, "y": 0.5, "h": 0.045},
            {"type": "chip", "text": "1", "x": 0.1, "y": 0.6, "d": 0.03},
            {"type": "hairline", "x": 0.05, "y": 0.7, "w": 0.9},
            {"type": "arrow", "x": 0.4, "y": 0.6, "w": 0.03},
            {"type": "sowhat", "text": "요약 한 줄입니다."},
        ],
    }
    (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))

    p2 = tmp_path / "pages" / "p02"
    p2.mkdir(parents=True)
    spec2 = {
        "page": 2, "variant": "dark",
        "items": [
            {"type": "glow", "cx": 0.1, "cy": 0.8, "sigma": 0.3},
            {"type": "serif_label", "text": "PROPOSAL 2026.", "x": 0.3, "y": 0.08, "w": 0.4},
            {"type": "headline", "x": 0.1, "y": 0.35, "w": 0.8, "align": "center",
             "runs": [[{"t": "지역 청년이 "}, {"t": "AI", "color": "E8622C"}, {"t": "로"}]]},
        ],
    }
    (p2 / "layout_spec.json").write_text(json.dumps(spec2, ensure_ascii=False))
    return tmp_path


class TestRenderDeck:
    def test_creates_pptx_with_two_slides(self, project):
        out = render_deck(project)
        prs = Presentation(out)
        assert len(prs.slides) == 2
        assert prs.slide_width == SLIDE_W and prs.slide_height == SLIDE_H

    def test_paper_page_has_native_shapes(self, project):
        out = render_deck(project)
        prs = Presentation(out)
        s1 = prs.slides[0]
        # 배경 + 라벨 + 헤드라인 + 카드 + 본문 + 필 + 칩 + 헤어라인 + 화살표 + sowhat(선+텍스트)
        assert len(s1.shapes) == 11

    def test_headline_font_is_paperlogy_bold(self, project):
        out = render_deck(project)
        prs = Presentation(out)
        s1 = prs.slides[0]
        texts = [sh for sh in s1.shapes if sh.has_text_frame
                 and "헤드라인" in sh.text_frame.text]
        run = texts[0].text_frame.paragraphs[0].runs[0]
        assert run.font.name == "Paperlogy 7 Bold"

    def test_dark_page_accent_run(self, project):
        out = render_deck(project)
        prs = Presentation(out)
        s2 = prs.slides[1]
        for sh in s2.shapes:
            if sh.has_text_frame and "AI" in sh.text_frame.text:
                runs = sh.text_frame.paragraphs[0].runs
                ai = [r for r in runs if r.text == "AI"][0]
                assert str(ai.font.color.rgb) == "E8622C"
                break
        else:
            pytest.fail("accent run not found")

    def test_pill_supports_grad_and_shadow(self, tmp_path):
        """pill도 card처럼 grad(그라데이션)·shadow(드롭섀도우)를 받는다."""
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "paper", "items": [
            {"type": "pill", "text": "공간 운영도", "x": 0.1, "y": 0.3, "h": 0.04,
             "w": 0.2, "grad": ["2E75B6", "1D4E86"], "shadow": True,
             "shadow_alpha": 16},
        ]}
        (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        out = render_deck(tmp_path)
        prs = Presentation(out)
        # 배경 + 필 = 2개. 필 도형 XML에 gradFill과 outerShdw가 있어야 한다.
        pill = [sh for sh in prs.slides[0].shapes if sh.has_text_frame
                and "공간 운영도" in sh.text_frame.text][0]
        xml = pill._element.xml
        assert "gradFill" in xml
        assert "outerShdw" in xml
        # 텍스트는 세로 중앙정렬(MIDDLE) 유지
        assert pill.text_frame.paragraphs[0].runs[0].text == "공간 운영도"

    def test_poly_and_text_rotation(self, tmp_path):
        """poly(프리폼 다각형)와 회전된 text(아이소메트릭 라벨) 지원."""
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "paper", "items": [
            {"type": "poly", "fill": "2C3E50",
             "points": [[0.2, 0.3], [0.5, 0.2], [0.6, 0.4], [0.3, 0.5]]},
            {"type": "text", "text": "세미나실", "x": 0.3, "y": 0.32, "w": 0.2,
             "size": 12, "color": "FFFFFF", "rot": -30},
        ]}
        (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        out = render_deck(tmp_path)
        prs = Presentation(out)
        shapes = list(prs.slides[0].shapes)
        # 프리폼 도형이 존재(custGeom XML)
        assert any("custGeom" in sh._element.xml for sh in shapes)
        # 회전 라벨
        lbl = [sh for sh in shapes if sh.has_text_frame
               and "세미나실" in sh.text_frame.text][0]
        assert abs((lbl.rotation % 360) - (-30 % 360)) < 0.01

    def test_bars_and_steps_primitives(self, tmp_path):
        import json as _json
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "paper", "items": [
            {"type": "bars", "x": 0.1, "y": 0.3, "w": 0.8, "h": 0.3,
             "items": [{"label": "창의성", "value": 30, "highlight": True},
                       {"label": "완성도", "value": 25}]},
            {"type": "steps", "x": 0.05, "y": 0.7, "w": 0.9,
             "items": ["기획", "제작", "발표"]},
        ]}
        (pdir / "layout_spec.json").write_text(_json.dumps(spec, ensure_ascii=False))
        out = render_deck(tmp_path)
        prs = Presentation(out)
        shapes = prs.slides[0].shapes
        # bg + bars(2행×3개) + steps(칩3+라벨3+화살표2)
        assert len(shapes) == 1 + 6 + 8

    def test_glow_asset_created_and_reused(self, project):
        render_deck(project)
        pngs = list((project / "assets").glob("glow_*.png"))
        assert len(pngs) == 1


class TestThemeRegistry:
    def test_json_theme_from_project_dir(self, tmp_path):
        import json as _json
        (tmp_path / "themes").mkdir()
        (tmp_path / "themes" / "mytheme.json").write_text(_json.dumps({
            "bg": "112233", "accent": "FF0000",
            "fonts": {"display_b": ["The Jamsil 5 Bold", False]},
        }))
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        (pdir / "layout_spec.json").write_text(_json.dumps({
            "page": 1, "theme": "mytheme", "variant": "paper",
            "items": [{"type": "headline", "lines": ["제목"], "x": 0.1, "y": 0.1, "w": 0.8}],
        }))
        out = render_deck(tmp_path)
        prs = Presentation(out)
        s = prs.slides[0]
        bg = s.shapes[0]
        assert str(bg.fill.fore_color.rgb) == "112233"
        run = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text][0] \
            .text_frame.paragraphs[0].runs[0]
        assert run.font.name == "The Jamsil 5 Bold"

    def test_inline_theme_tokens(self, tmp_path):
        import json as _json
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        (pdir / "layout_spec.json").write_text(_json.dumps({
            "page": 1, "theme": "oneoff", "variant": "paper",
            "theme_tokens": {"bg": "ABCDEF"},
            "items": [],
        }))
        out = render_deck(tmp_path)
        prs = Presentation(out)
        assert str(prs.slides[0].shapes[0].fill.fore_color.rgb) == "ABCDEF"

    def test_unknown_theme_raises(self, tmp_path):
        import json as _json
        import pytest as _pytest
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        (pdir / "layout_spec.json").write_text(_json.dumps({
            "page": 1, "theme": "nope", "items": []}))
        with _pytest.raises(KeyError):
            render_deck(tmp_path)


class TestInfographicPrimitives:
    def _render(self, tmp_path, items):
        import json as _json
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        (pdir / "layout_spec.json").write_text(_json.dumps(
            {"page": 1, "theme": "gyoan", "variant": "paper", "items": items},
            ensure_ascii=False))
        return Presentation(render_deck(tmp_path))

    def test_shape_hexagon_and_pie(self, tmp_path):
        prs = self._render(tmp_path, [
            {"type": "shape", "shape": "hexagon", "x": 0.1, "y": 0.1, "w": 0.1, "h": 0.1,
             "fill": "112233", "text": "01"},
            {"type": "shape", "shape": "pie", "x": 0.3, "y": 0.1, "w": 0.15, "h": 0.15,
             "fill": "E8622C", "adj": [0, 120]},
            {"type": "shape", "shape": "parallelogram", "x": 0.5, "y": 0.1, "w": 0.2, "h": 0.06,
             "fill": "none", "outline": "112233"},
        ])
        # bg + hexagon + pie + parallelogram
        assert len(prs.slides[0].shapes) == 4

    def test_card_gradient_and_shadow(self, tmp_path):
        prs = self._render(tmp_path, [
            {"type": "card", "x": 0.1, "y": 0.1, "w": 0.5, "h": 0.1,
             "grad": ["26467B", "1F3864"], "shadow": True},
        ])
        card = prs.slides[0].shapes[1]
        xml = card._element.xml
        assert "gradFill" in xml and "26467B" in xml
        assert "outerShdw" in xml

    def test_hairline_dash(self, tmp_path):
        prs = self._render(tmp_path, [
            {"type": "hairline", "x": 0.1, "y": 0.5, "w": 0.1, "dash": True},
        ])
        assert len(prs.slides[0].shapes) > 5  # 세그먼트 다수

    def test_card_outline_dash(self, tmp_path):
        """빈 슬롯(사진 자리) 선언용 — card outline_dash로 파선 테두리."""
        prs = self._render(tmp_path, [
            {"type": "card", "x": 0.1, "y": 0.1, "w": 0.2, "h": 0.1,
             "fill": "F3EFE7", "outline": "C8912F", "outline_dash": True,
             "outline_w": 0.9},
        ])
        card = prs.slides[0].shapes[1]
        xml = card._element.xml
        assert 'prstDash' in xml and 'val="dash"' in xml
        assert "C8912F" in xml


class TestPageSize:
    def _proj(self, tmp_path, page_size=None, items=None):
        import json as _json
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "theme": "gyoan", "variant": "paper",
                "items": items or [{"type": "headline", "lines": ["제목"],
                                    "x": 0.1, "y": 0.1, "w": 0.8}]}
        if page_size:
            spec["page_size"] = page_size
        (pdir / "layout_spec.json").write_text(_json.dumps(spec))
        return tmp_path

    def test_default_is_16_9(self, tmp_path):
        out = render_deck(self._proj(tmp_path))
        prs = Presentation(out)
        assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)

    def test_a4_portrait(self, tmp_path):
        out = render_deck(self._proj(tmp_path, page_size="a4p"))
        prs = Presentation(out)
        assert (prs.slide_width, prs.slide_height) == (6858000, 9906000)
        # 배경 사각형이 전체 페이지를 덮는다
        bg = prs.slides[0].shapes[0]
        assert (bg.width, bg.height) == (6858000, 9906000)

    def test_a4_portrait_glow_aspect(self, tmp_path):
        from PIL import Image as _Image
        proj = self._proj(tmp_path, page_size="a4p", items=[
            {"type": "glow", "cx": 0.5, "cy": 0.8, "intensity": 0.5}])
        render_deck(proj)
        pngs = list((proj / "assets").glob("glow_*.png"))
        assert pngs, "glow png 생성돼야 함"
        w, h = _Image.open(pngs[0]).size
        assert h > w  # 세로 종횡비

    def test_poster_3x4(self, tmp_path):
        """세로 3:4 포스터(키비주얼·인스타·에타 마스터)."""
        out = render_deck(self._proj(tmp_path, page_size="3:4"))
        prs = Presentation(out)
        assert (prs.slide_width, prs.slide_height) == (6858000, 9144000)
        bg = prs.slides[0].shapes[0]
        assert (bg.width, bg.height) == (6858000, 9144000)

    def test_page_size_restored_between_decks(self, tmp_path):
        # a4p 덱 렌더 후 기본 덱이 다시 16:9로 렌더되는지 (전역 오염 방지)
        out1 = render_deck(self._proj(tmp_path / "a", page_size="a4p"))
        out2 = render_deck(self._proj(tmp_path / "b"))
        assert Presentation(out2).slide_width == 12192000


class TestBandPrimitive:
    """채워진 바 안 라벨+본문을 세로·가로 자동 중앙 정렬하는 band 프리미티브."""

    def _band_proj(self, tmp_path, item):
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "paper", "items": [item]}
        (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        return tmp_path

    def _band_shape(self, prs):
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        for sh in prs.slides[0].shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                return sh
        return None

    def test_label_band_vertically_and_horizontally_centered(self, tmp_path):
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        proj = self._band_proj(tmp_path, {
            "type": "band", "x": 0.055, "y": 0.845, "w": 0.89, "h": 0.062,
            "fill": "1C2534", "label": "SO WHAT", "label_color": "C8974E",
            "text": "본문 한 줄입니다.", "text_color": "F3F0E9"})
        prs = Presentation(render_deck(proj))
        band = self._band_shape(prs)
        assert band is not None, "band 도형이 텍스트와 함께 생성돼야 함"
        # 세로 중앙: 손좌표가 아니라 텍스트프레임 앵커로 보장
        assert band.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        # 라벨+본문이 한 문단에 같이(사이 빈공백 없음), 가로 중앙 정렬
        assert band.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
        txt = band.text_frame.text
        assert "SO WHAT" in txt and "본문 한 줄입니다." in txt

    def test_labelless_band_centers_text(self, tmp_path):
        from pptx.enum.text import PP_ALIGN
        proj = self._band_proj(tmp_path, {
            "type": "band", "x": 0.28, "y": 0.85, "w": 0.44, "h": 0.072,
            "fill": "1C2534", "text": "다음 날 아침이, 고요를 증명합니다.", "align": "center"})
        prs = Presentation(render_deck(proj))
        band = self._band_shape(prs)
        assert band is not None
        assert band.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
        assert "고요를 증명합니다." in band.text_frame.text


class TestScrimPrimitive:
    """풀블리드 이미지용 알파 그라데이션 스크림."""

    def test_scrim_creates_rgba_alpha_gradient(self, tmp_path):
        from PIL import Image as _Image
        import numpy as _np
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "dark", "items": [
            {"type": "scrim", "edge": "bottom", "h": 0.5, "color": "000000", "strength": 0.9}]}
        (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        render_deck(tmp_path)
        pngs = list((tmp_path / "assets").glob("scrim_*.png"))
        assert pngs, "scrim png 생성돼야 함"
        im = _Image.open(pngs[0])
        assert im.mode == "RGBA"
        a = _np.asarray(im)[..., 3]
        # 안쪽(위)은 투명, 바깥(아래)은 진함 — 세로 알파 그라데이션
        assert a[0].mean() < 20 and a[-1].mean() > 180


class TestNativeChartTable:
    """그래프·표는 PPT 네이티브 개체(편집 가능한 데이터)로 — card 사각형 금지."""

    def _slide(self, tmp_path, items):
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        spec = {"page": 1, "variant": "paper", "items": items}
        (pdir / "layout_spec.json").write_text(json.dumps(spec, ensure_ascii=False))
        return Presentation(render_deck(tmp_path)).slides[0]

    def test_chart_is_native_graphicframe_with_editable_data(self, tmp_path):
        s = self._slide(tmp_path, [
            {"type": "chart", "kind": "bar", "x": 0.1, "y": 0.3, "w": 0.5, "h": 0.4,
             "categories": ["국내주식", "해외주식", "채권"],
             "series": [{"name": "배분", "values": [25, 30, 20]}],
             "colors": ["2A3B57", "BE9040", "5B6B84"], "value_labels": True},
        ])
        gfs = [sh for sh in s.shapes if sh.has_chart]
        assert len(gfs) == 1, "네이티브 차트 GraphicFrame 하나"
        chart = gfs[0].chart
        # 카테고리·값이 실제 편집 가능한 차트 데이터로 들어감
        assert list(chart.plots[0].categories) == ["국내주식", "해외주식", "채권"]
        assert list(chart.series[0].values) == [25.0, 30.0, 20.0]
        # 한글이 깨지지 않도록 ea 폰트가 차트 txPr에 주입됨
        assert "<a:ea" in chart.font._rPr.xml
        # 포인트별 색상 적용(단일 시리즈) — 차트 파트 XML에 dPt 색상
        assert "BE9040" in chart._chartSpace.xml

    def test_table_is_native_with_cell_text_and_ea_font(self, tmp_path):
        s = self._slide(tmp_path, [
            {"type": "table", "x": 0.1, "y": 0.3, "w": 0.8, "h": 0.3,
             "rows": [["항목", "비중"], ["국내주식", "25%"], ["해외주식", "30%"]],
             "header": True, "col_w": [0.6, 0.4], "align": ["left", "right"]},
        ])
        tbls = [sh for sh in s.shapes if sh.has_table]
        assert len(tbls) == 1, "네이티브 표 하나"
        table = tbls[0].table
        assert len(table.rows) == 3 and len(table.columns) == 2
        assert table.cell(0, 0).text == "항목"
        assert table.cell(2, 1).text == "30%"
        # 한글 셀에 ea 폰트 주입
        run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
        assert "<a:ea" in run._r.xml

    def test_num_xl_default_line_spacing_is_tight(self, tmp_path):
        """빅넘버는 기본 줄간격 1.0 — 밑줄/헤어라인과 겹치지 않게."""
        s = self._slide(tmp_path, [
            {"type": "text", "text": "23분", "x": 0.1, "y": 0.3, "w": 0.2, "style": "num_xl"},
        ])
        num = [sh for sh in s.shapes if sh.has_text_frame and "23분" in sh.text_frame.text][0]
        assert num.text_frame.paragraphs[0].line_spacing == 1.0


# ── serif_label: align/color/size 오버라이드 (2026-07-21 버그) ─────────────
class TestSerifLabelOverrides:
    """serif_label이 align을 계산만 하고 버려 좌마진 라벨이 가운데로 밀렸다."""

    def _render(self, tmp_path, item):
        from pptx import Presentation
        from pptx.util import Emu
        import pptm.design.gyoan as G
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(G.SLIDE_W), Emu(G.SLIDE_H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        G.render_page(slide, {"page": 1, "theme": "gyoan", "items": [item]}, tmp_path)
        return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text][-1]

    def test_align_left를_존중한다(self, tmp_path):
        from pptx.enum.text import PP_ALIGN
        shp = self._render(tmp_path, {
            "type": "serif_label", "x": 0.07, "y": 0.1, "w": 0.5,
            "text": "潭 · DAM", "align": "left"})
        assert shp.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT

    def test_기본은_여전히_center다(self, tmp_path):
        from pptx.enum.text import PP_ALIGN
        shp = self._render(tmp_path, {
            "type": "serif_label", "x": 0.07, "y": 0.1, "w": 0.5, "text": "DAM"})
        assert shp.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER

    def test_color를_존중한다(self, tmp_path):
        shp = self._render(tmp_path, {
            "type": "serif_label", "x": 0.07, "y": 0.1, "w": 0.5,
            "text": "DAM", "color": "C9A96A"})
        assert str(shp.text_frame.paragraphs[0].runs[0].font.color.rgb) == "C9A96A"

    def test_size를_존중한다(self, tmp_path):
        shp = self._render(tmp_path, {
            "type": "serif_label", "x": 0.07, "y": 0.1, "w": 0.5,
            "text": "DAM", "size": 12})
        assert shp.text_frame.paragraphs[0].runs[0].font.size.pt == 12

    def test_이탤릭은_유지된다(self, tmp_path):
        shp = self._render(tmp_path, {
            "type": "serif_label", "x": 0.07, "y": 0.1, "w": 0.5, "text": "DAM"})
        assert shp.text_frame.paragraphs[0].runs[0].font.italic is True


class TestTracking:
    """자간(letter-spacing) — 매거진 마스트헤드용 `tracking` 키.

    OOXML은 rPr/@spc(1/100 pt)로 자간을 준다. python-pptx가 노출하지 않아
    직접 세팅하므로 회귀를 테스트로 묶는다.
    """

    def _spec(self, tmp_path, items):
        pdir = tmp_path / "pages" / "p01"
        pdir.mkdir(parents=True)
        (pdir / "layout_spec.json").write_text(
            json.dumps({"page": 1, "items": items}, ensure_ascii=False))
        return render_deck(tmp_path)

    def _runs(self, out, needle):
        prs = Presentation(out)
        for sh in prs.slides[0].shapes:
            if sh.has_text_frame and needle in sh.text_frame.text:
                return sh.text_frame.paragraphs[0].runs
        raise AssertionError(f"'{needle}' 텍스트 없음")

    def test_아이템_tracking이_spc로_들어간다(self, tmp_path):
        out = self._spec(tmp_path, [
            {"type": "text", "text": "2026 JOB FAIR", "x": 0.08, "y": 0.1,
             "w": 0.6, "tracking": 2.5},
        ])
        run = self._runs(out, "JOB FAIR")[0]
        assert run._r.get_or_add_rPr().get("spc") == "250"

    def test_tracking_없으면_spc도_없다(self, tmp_path):
        out = self._spec(tmp_path, [
            {"type": "text", "text": "그냥 본문", "x": 0.08, "y": 0.1, "w": 0.6},
        ])
        run = self._runs(out, "그냥 본문")[0]
        assert run._r.get_or_add_rPr().get("spc") is None

    def test_음수_tracking도_지원한다(self, tmp_path):
        out = self._spec(tmp_path, [
            {"type": "text", "text": "빽빽하게", "x": 0.08, "y": 0.1,
             "w": 0.6, "tracking": -0.4},
        ])
        run = self._runs(out, "빽빽하게")[0]
        assert run._r.get_or_add_rPr().get("spc") == "-40"

    def test_런별_tracking이_아이템값을_덮는다(self, tmp_path):
        out = self._spec(tmp_path, [
            {"type": "text", "x": 0.08, "y": 0.1, "w": 0.6, "tracking": 2.0,
             "runs": [[{"t": "AA"}, {"t": "BB", "tracking": 0.5}]]},
        ])
        runs = self._runs(out, "AABB")
        assert runs[0]._r.get_or_add_rPr().get("spc") == "200"
        assert runs[1]._r.get_or_add_rPr().get("spc") == "50"

    def test_headline과_label에도_먹는다(self, tmp_path):
        out = self._spec(tmp_path, [
            {"type": "headline", "lines": ["헤드입니다."], "x": 0.08, "y": 0.2,
             "w": 0.8, "tracking": 1.0},
            {"type": "label", "text": "라벨", "x": 0.08, "y": 0.1, "tracking": 3.0},
        ])
        assert self._runs(out, "헤드입니다.")[0]._r.get_or_add_rPr().get("spc") == "100"
        assert self._runs(out, "라벨")[0]._r.get_or_add_rPr().get("spc") == "300"
