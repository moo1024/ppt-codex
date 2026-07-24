"""export 스테이지 테스트 — soffice/pdftoppm 모킹."""
import io
import json
from pathlib import Path
from unittest.mock import patch, MagicMock

import pytest
from PIL import Image

from pptm.stages.export import run_export, export_thumbnails


def _make_fake_pptx(path: Path) -> None:
    """python-pptx로 진짜 빈 pptx 생성."""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))


def _make_fake_pdf(path: Path) -> None:
    path.write_bytes(b"%PDF-1.4\n1 0 obj\n<< >>\nendobj\nxref\ntrailer\n<< >>\n%%EOF\n")


def _make_fake_thumb(path: Path, w=320, h=180) -> None:
    img = Image.new("RGB", (w, h), color=(30, 30, 30))
    img.save(str(path), format="JPEG")


class TestRunExport:
    def _setup(self, tmp_path: Path) -> Path:
        proj = tmp_path / "demo"
        out_dir = proj / "output"
        out_dir.mkdir(parents=True)
        pptx_path = out_dir / "deck.pptx"
        _make_fake_pptx(pptx_path)
        return proj

    def test_creates_pdf(self, tmp_path):
        proj = self._setup(tmp_path)
        out_dir = proj / "output"

        def fake_soffice(*args, **kwargs):
            m = MagicMock()
            m.returncode = 0
            cmd = args[0]
            if "--convert-to" in cmd:
                fmt = cmd[cmd.index("--convert-to") + 1]
                out_arg = cmd[cmd.index("--outdir") + 1]
                src = Path(cmd[-1])
                if "pdf" in fmt:
                    # fallback 경로: --outdir outdir로 직접 PDF 생성
                    (Path(out_arg) / (src.stem + ".pdf")).write_bytes(b"%PDF-fake")
                # ODP는 생성 안 함 → fallback 경로 진입
            return m

        with patch("pptm.stages.export.subprocess.run", side_effect=fake_soffice):
            run_export(project_dir=proj)

        assert (out_dir / "deck.pdf").exists()

    def test_creates_thumbnails_dir(self, tmp_path):
        proj = self._setup(tmp_path)
        out_dir = proj / "output"
        _make_fake_pdf(out_dir / "deck.pdf")

        def fake_pdftoppm(*args, **kwargs):
            m = MagicMock()
            m.returncode = 0
            # pdftoppm 호출 흉내 — slide-1.jpg 생성
            cmd = args[0]
            prefix = Path(cmd[-1])
            thumb_dir = prefix.parent
            _make_fake_thumb(thumb_dir / "slide-1.jpg")
            return m

        with patch("pptm.stages.export.subprocess.run", side_effect=fake_pdftoppm):
            thumbs = export_thumbnails(out_dir / "deck.pdf", out_dir, dpi=72)

        assert (out_dir / "thumbnails").is_dir()
        assert any("slide" in t.name for t in thumbs)

    def test_skips_if_no_pptx(self, tmp_path):
        proj = tmp_path / "demo"
        (proj / "output").mkdir(parents=True)
        # deck.pptx 없음 → 조용히 반환
        run_export(project_dir=proj)  # 예외 없어야 함

    def test_grid_created_when_multiple_thumbs(self, tmp_path):
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        thumb_dir = out_dir / "thumbnails"
        thumb_dir.mkdir()

        # 가짜 썸네일 4장 생성
        fake_thumbs = []
        for i in range(1, 5):
            p = thumb_dir / f"slide-{i:02d}.jpg"
            _make_fake_thumb(p)
            fake_thumbs.append(p)

        from pptm.stages.export import _make_grids
        grids = _make_grids(fake_thumbs, thumb_dir)
        assert len(grids) >= 1
        assert all(g.suffix == ".jpg" for g in grids)
